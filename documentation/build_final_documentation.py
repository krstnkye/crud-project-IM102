from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUTPUT = "Product_Inventory_System_Documentation.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, header_fill="D9EAF7"):
    table.style = "Table Grid"
    table.autofit = True
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(10)
            if r == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    if level == 1:
        paragraph.style = doc.styles["Heading 1"]
    elif level == 2:
        paragraph.style = doc.styles["Heading 2"]
    else:
        paragraph.style = doc.styles["Heading 3"]
    run = paragraph.add_run(text)
    run.font.color.rgb = None
    return paragraph


def add_para(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.08
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbered(doc, items):
    for index, item in enumerate(items, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run(f"{index}. {item}")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = str(value)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    style_table(table)
    doc.add_paragraph()
    return table


def configure_document(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for style_name, size in (("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Product Inventory System Documentation")
    run.bold = True
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("IM101 CRUD Operations Project").bold = True

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run("Prepared by: ______________________________\n")
    info.add_run("Instructor: ______________________________\n")
    info.add_run("Date: October 8, 2026")

    doc.add_page_break()

    add_heading(doc, "Overview of the Proposed Software")
    add_para(
        doc,
        "The proposed software is a Product Inventory System. It is a simple desktop application that helps a user keep a list of products. The user can add a product, view saved products, update product details, and delete products that are no longer needed.",
    )
    add_para(
        doc,
        "The system was made for a beginner CRUD project. It focuses on one product list, clear buttons, automatic totals, and saved records. It uses Python for the program, Tkinter for the window design, and SQLite for saving product information.",
    )

    add_heading(doc, "Description of the Software")
    add_para(
        doc,
        "The Product Inventory System opens as a desktop window. At the top, it shows total products, total stock, low stock products, and total value. The left side contains the product form. The right side contains the product list where saved records are shown.",
    )
    add_para(
        doc,
        "The application is useful for small inventory records such as school supplies, food items, drinks, cleaning materials, electronics, and other simple products. The records remain saved because the program uses an SQLite database file named inventory.db.",
    )

    add_heading(doc, "Features of the Software")
    add_table(
        doc,
        ["Feature", "Simple Description"],
        [
            ["Add Product", "The user can enter product details and save them."],
            ["View Products", "The user can see all saved products in the product list."],
            ["Update Product", "The user can select a product and change its details."],
            ["Delete Product", "The user can remove a selected product after confirmation."],
            ["Clear Fields", "The user can empty the form before typing new details."],
            ["Dashboard Totals", "The system automatically shows total products, stock, low stock, and total value."],
        ],
        [1.8, 4.9],
    )

    add_heading(doc, "Software Development Team Structure")
    add_table(
        doc,
        ["Member", "Assigned Task"],
        [
            ["Member 1", "Window design and form layout"],
            ["Member 2", "Database and saving of records"],
            ["Member 3", "Checking inputs and testing"],
            ["Member 4", "Documentation, screenshots, and report"],
        ],
        [1.4, 5.3],
    )

    add_heading(doc, "Software Development Tool Choice of SDLC Model")
    add_para(
        doc,
        "The team may use the waterfall model for this project because the system is small and easy to plan step by step. The project starts with planning, then analysis, design, coding, testing, and documentation.",
    )
    add_table(
        doc,
        ["Phase", "What Was Done"],
        [
            ["Planning", "The team chose an inventory system as the project topic."],
            ["Analysis", "The team listed the needed product details and actions."],
            ["Design", "The team planned the window, form, buttons, and product list."],
            ["Coding", "The team created the Python, Tkinter, and SQLite files."],
            ["Testing", "The team checked adding, viewing, updating, and deleting records."],
            ["Documentation", "The team prepared the written report and user guide."],
        ],
        [1.4, 5.3],
    )

    add_heading(doc, "Software Planning")
    add_para(
        doc,
        "The project was planned as a simple inventory program that can be completed within the class project period. The first task was to identify the product fields. The next task was to design the form and product list. After that, the team created the database and tested each button.",
    )
    add_table(
        doc,
        ["Week", "Planned Activity"],
        [
            ["Week 1", "Choose project topic and list needed features."],
            ["Week 2", "Prepare the product fields and database design."],
            ["Week 3", "Create the window layout and product form."],
            ["Week 4", "Add the save, update, delete, and clear actions."],
            ["Week 5", "Test the program and fix errors."],
            ["Week 6", "Finish documentation and screenshots."],
        ],
        [1.2, 5.5],
    )

    add_heading(doc, "Software Analysis")
    add_para(
        doc,
        "The main problem is that product records can be hard to manage when written only on paper or saved in separate notes. The Product Inventory System solves this by keeping records in one application. It also updates totals automatically so the user can quickly see the number of products and stock available.",
    )

    add_heading(doc, "Software Specification")
    add_table(
        doc,
        ["Item", "Specification"],
        [
            ["Program Name", "Product Inventory System"],
            ["Programming Language", "Python"],
            ["Window Library", "Tkinter"],
            ["Saved File", "inventory.db"],
            ["Database", "SQLite"],
            ["Main File", "main.py"],
            ["Database File", "database.py"],
        ],
        [2.0, 4.7],
    )

    add_heading(doc, "Hardware Specification")
    add_table(
        doc,
        ["Requirement", "Minimum Specification"],
        [
            ["Computer", "Laptop or desktop computer"],
            ["Operating System", "Windows 10 or Windows 11"],
            ["Memory", "At least 4 GB RAM"],
            ["Storage", "At least 100 MB free space"],
            ["Other", "Keyboard and mouse"],
        ],
        [2.0, 4.7],
    )

    add_heading(doc, "Functional Requirements")
    add_table(
        doc,
        ["No.", "Requirement"],
        [
            ["FR 1", "The system shall allow the user to add a product."],
            ["FR 2", "The system shall show all saved products in a product list."],
            ["FR 3", "The system shall allow the user to update a selected product."],
            ["FR 4", "The system shall allow the user to delete a selected product."],
            ["FR 5", "The system shall update the totals after every change."],
            ["FR 6", "The system shall show messages for successful actions and wrong inputs."],
        ],
        [1.0, 5.7],
    )

    add_heading(doc, "Non Functional Requirements")
    add_table(
        doc,
        ["Requirement", "Description"],
        [
            ["Ease of Use", "The buttons and labels should be easy to understand."],
            ["Reliability", "The records should remain saved after closing the program."],
            ["Speed", "The product list and totals should refresh quickly."],
            ["Safety", "The system should ask first before deleting a product."],
            ["Maintainability", "The files should be organized and easy to edit."],
        ],
        [1.8, 4.9],
    )

    add_heading(doc, "Software Designing")
    add_para(
        doc,
        "The design follows a simple dashboard layout. The upper part displays the totals. The left part has the product form. The right part has the product list. This layout helps the user see the form and records at the same time.",
    )
    add_table(
        doc,
        ["Part of the Window", "Purpose"],
        [
            ["Top Summary", "Shows total products, total stock, low stock, and total value."],
            ["Product Form", "Allows the user to type product name, category, quantity, and price."],
            ["Product List", "Shows saved product records."],
            ["Buttons", "Allows the user to add, update, delete, and clear fields."],
        ],
        [2.0, 4.7],
    )

    add_heading(doc, "Database Design")
    add_para(doc, "The project uses one table named products. This table stores all product records.")
    add_table(
        doc,
        ["Field", "Data Type", "Description"],
        [
            ["product_id", "INTEGER", "Automatic product number"],
            ["product_name", "TEXT", "Name of the product"],
            ["category", "TEXT", "Group of the product"],
            ["quantity", "INTEGER", "Number of pieces available"],
            ["unit_price", "REAL", "Price of one piece"],
        ],
        [1.6, 1.4, 3.7],
    )

    add_heading(doc, "Role Activity Diagram")
    add_para(doc, "The user follows this basic process when using the system:")
    add_numbered(
        doc,
        [
            "Open the Product Inventory System.",
            "View the product list and totals.",
            "Enter product details in the form.",
            "Click Add Product to save a new product.",
            "Select a product from the list to update or delete it.",
            "Check if the product list and totals changed correctly.",
            "Close the program when finished.",
        ],
    )

    add_heading(doc, "Software Requirements")
    add_para(
        doc,
        "The system needs Python installed on the computer. Tkinter is used for the window, buttons, labels, forms, and pop-up messages. SQLite is used for the saved product file.",
    )
    add_bullets(
        doc,
        [
            "Python is used for the program instructions.",
            "Tkinter is used for the desktop window.",
            "SQLite is used to save products in inventory.db.",
            "VS Code or IDLE may be used to open and run the project.",
        ],
    )

    add_heading(doc, "Software Implementation")
    add_para(
        doc,
        "The system was implemented using two main Python files. The main.py file contains the window design and button actions. The database.py file contains the saving, reading, updating, deleting, and total counting functions.",
    )
    add_table(
        doc,
        ["File", "Purpose"],
        [
            ["main.py", "Shows the window and handles button clicks."],
            ["database.py", "Creates the database and handles product records."],
            ["inventory.db", "Stores the saved products."],
            ["README.md", "Gives simple instructions for running the project."],
        ],
        [1.6, 5.1],
    )

    add_heading(doc, "Installation Flow Procedure")
    add_numbered(
        doc,
        [
            "Install Python on the computer.",
            "Open the InventoryManagement folder.",
            "Open Command Prompt or VS Code terminal in the folder.",
            "Type python main.py and press Enter.",
            "Wait for the Product Inventory System window to open.",
        ],
    )

    add_heading(doc, "User Manual")
    add_heading(doc, "Adding a Product", level=2)
    add_numbered(
        doc,
        [
            "Type the product name.",
            "Choose or type the category.",
            "Type the quantity.",
            "Type the price each.",
            "Click Add Product.",
            "Check if the product appears in the product list.",
        ],
    )
    add_heading(doc, "Updating a Product", level=2)
    add_numbered(
        doc,
        [
            "Click a product from the product list.",
            "Change the details in the form.",
            "Click Update.",
            "Check if the product list shows the new details.",
        ],
    )
    add_heading(doc, "Deleting a Product", level=2)
    add_numbered(
        doc,
        [
            "Click a product from the product list.",
            "Click Delete.",
            "Choose Yes when the confirmation message appears.",
            "Check if the product was removed from the list.",
        ],
    )

    add_heading(doc, "Software Testing and Integration")
    add_para(
        doc,
        "The system should be tested by using both correct and wrong inputs. This helps check if the program works properly and if the messages are clear.",
    )
    add_table(
        doc,
        ["Test ID", "Test Condition", "Expected Result", "Actual Result", "Remarks"],
        [
            ["TC01", "Add a valid product", "Product is saved and shown in the list", "To be recorded", "Pending"],
            ["TC02", "Leave product name blank", "Message asks for product name", "To be recorded", "Pending"],
            ["TC03", "Type letters in quantity", "Message says quantity must be a whole number", "To be recorded", "Pending"],
            ["TC04", "Type a negative quantity", "Message says quantity cannot be negative", "To be recorded", "Pending"],
            ["TC05", "Update a selected product", "Product details are changed", "To be recorded", "Pending"],
            ["TC06", "Delete a selected product", "Product is removed after confirmation", "To be recorded", "Pending"],
            ["TC07", "Close and reopen the program", "Saved products are still shown", "To be recorded", "Pending"],
        ],
        [0.7, 1.55, 2.0, 1.25, 0.9],
    )

    add_heading(doc, "Error Handling")
    add_table(
        doc,
        ["Situation", "Message Shown"],
        [
            ["Empty product name", "Please enter a product name."],
            ["Quantity has letters", "Quantity must be a whole number."],
            ["Negative quantity", "Quantity cannot be negative."],
            ["Invalid price", "Please enter a valid price."],
            ["Update without choosing a product", "Please select a product to update."],
            ["Delete without choosing a product", "Please select a product to delete."],
            ["Delete confirmation", "Are you sure you want to delete this product?"],
        ],
        [2.4, 4.3],
    )

    add_heading(doc, "Results and Discussion")
    add_para(
        doc,
        "The Product Inventory System was able to demonstrate the four main CRUD actions. The user can add a product, read or view saved products, update selected product details, and delete selected products. The dashboard totals also update after changes are made.",
    )
    add_para(
        doc,
        "The system is suitable for a beginner project because it uses simple product fields and one saved product table. The project also shows the importance of checking user input before saving records.",
    )

    add_heading(doc, "Summary Conclusion and Recommendation")
    add_heading(doc, "Summary", level=2)
    add_para(
        doc,
        "This project is a desktop inventory program that keeps product records in a saved file. It uses Python, Tkinter, and SQLite to show a form, product list, buttons, messages, and totals.",
    )
    add_heading(doc, "Conclusion", level=2)
    add_para(
        doc,
        "The project helped demonstrate how a simple information system works. It showed how to create a window, accept user input, save records, display records, change records, delete records, and keep data available after closing the program.",
    )
    add_heading(doc, "Recommendation", level=2)
    add_para(
        doc,
        "Future improvements may include a search bar, product categories report, low stock warning, printable inventory report, login screen, and export to Excel or CSV.",
    )

    add_heading(doc, "Appendix A Screenshots")
    add_para(
        doc,
        "Insert project screenshots in this section after running the program. Suggested screenshots are dashboard, add product, update product, delete product, and error message.",
    )
    add_table(
        doc,
        ["Screenshot Needed", "Description"],
        [
            ["Dashboard", "Shows the main window and totals."],
            ["Add Product", "Shows a product being added."],
            ["Update Product", "Shows a selected product being changed."],
            ["Delete Product", "Shows the delete confirmation message."],
            ["Error Message", "Shows an invalid input message."],
        ],
        [2.0, 4.7],
    )

    add_heading(doc, "Appendix B Project Folder")
    add_para(doc, "The project folder is named InventoryManagement.")
    add_table(
        doc,
        ["Folder or File", "Description"],
        [
            ["main.py", "Main program file"],
            ["database.py", "Database functions file"],
            ["inventory.db", "Saved product records"],
            ["documentation", "Report and flowchart files"],
            ["screenshots", "Folder for test screenshots"],
        ],
        [2.0, 4.7],
    )

    section = doc.add_section(WD_SECTION.CONTINUOUS)
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not footer.text:
            footer.add_run("Product Inventory System Documentation")

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
